"""Regenerate fixtures/condition-catalog.js from the Business Dictionary.

The workbook is treated as data. A row becomes an approved pricing
condition only when it has a Premium Category and an Identifier Name,
its "Relevant For:" contains Pricing, and its Pricing Condition Group is
populated. The "New Condition Requests" sheet holds pending requests and
is never read.

Usage:
    python3 tools/build-condition-catalog.py "/path/to/Business Dictionary Service.xlsx"
"""

import json
import os
import re
import sys

import openpyxl

SHEET = "Business Conditions"
OUT = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                   "fixtures", "condition-catalog.js")

# Column positions in the workbook, zero-based.
COL_SYS_ID = 0      # A  Systematic Identifier
COL_CATEGORY = 1    # B  Premium Category  (shown to users as "Category")
COL_NAME = 2        # C  Identifier Name
COL_DESCRIPTION = 3 # D  Description
COL_RELEVANT = 5    # F  Relevant For:
COL_GROUP = 6       # G  Pricing Condition Group


def norm(value):
    return re.sub(r"\s+", " ", str(value)).strip() if value is not None else ""


def tokens(value):
    """Split a multi-value cell. Separators are inconsistent in the source,
    so both semicolons and commas are accepted."""
    return [t.strip().lower() for t in re.split(r"[;,]", norm(value)) if t.strip()]


def build(path):
    book = openpyxl.load_workbook(path, data_only=True, read_only=True)
    if SHEET not in book.sheetnames:
        raise SystemExit(f"workbook has no {SHEET!r} sheet: {book.sheetnames}")

    order, by_category = [], {}
    for raw in book[SHEET].iter_rows(min_row=2, values_only=True):
        cells = list(raw) + [None] * 8
        category = norm(cells[COL_CATEGORY])
        name = norm(cells[COL_NAME])
        if not category or not name:
            continue
        if "pricing" not in tokens(cells[COL_RELEVANT]):
            continue
        groups = tokens(cells[COL_GROUP])
        if not groups:
            continue
        if category not in by_category:
            by_category[category] = {}
            order.append(category)
        key = name.lower()
        if key in by_category[category]:
            # Same condition listed twice: union the eligibility rather
            # than emitting a duplicate option.
            existing = by_category[category][key]
            existing["groups"] = sorted(set(existing["groups"]) | set(groups))
            continue
        by_category[category][key] = {
            "name": name,
            "description": norm(cells[COL_DESCRIPTION]),
            "sysId": norm(cells[COL_SYS_ID]),
            "groups": sorted(set(groups)),
        }

    rows = []
    for category in order:
        entries = sorted(by_category[category].values(),
                         key=lambda e: e["name"].lower())
        for entry in entries:
            scope = ("L" if "base rate" in entry["groups"] else "") \
                  + ("P" if "premium" in entry["groups"] else "")
            rows.append([category, entry["name"], entry["description"],
                         scope, entry["sysId"]])
    return rows


def render(rows):
    counts = {
        "total": len(rows),
        "line": sum(1 for r in rows if "L" in r[3]),
        "premium": sum(1 for r in rows if "P" in r[3]),
        "both": sum(1 for r in rows if r[3] == "LP"),
    }
    body = "\n".join("    " + json.dumps(r, ensure_ascii=False) + ","
                     for r in rows)
    return f'''/* Approved pricing conditions, generated from the Business Dictionary
 * Service workbook, sheet "Business Conditions".
 *
 * A row is included only when it has a Premium Category and an
 * Identifier Name, its "Relevant For:" contains Pricing, and its
 * Pricing Condition Group is populated. Rows from the "New Condition
 * Requests" sheet are pending, not approved, and are never included.
 *
 * Each entry is [category, name, description, scope, systematicId].
 * Scope is L for Base Rate, P for Premium, LP for both. The workbook
 * has no Systematic Identifier values yet, so the canonical key is a
 * normalized category + name composite; see conditionKey() in v2.js.
 *
 * Regenerate with tools/build-condition-catalog.py. Do not hand edit.
 */
window.RCMConditionCatalog = (function () {{
  "use strict";

  var ROWS = [
{body}
  ];

  return {{
    rows: ROWS,
    /* Source counts at generation time, asserted by qa_conditions.py. */
    counts: {json.dumps(counts)}
  }};
}}());
'''


if __name__ == "__main__":
    if len(sys.argv) < 2:
        raise SystemExit(__doc__)
    built = build(sys.argv[1])
    with open(OUT, "w", encoding="utf-8") as handle:
        handle.write(render(built))
    print(f"wrote {OUT}: {len(built)} conditions "
          f"({sum(1 for r in built if 'L' in r[3])} line, "
          f"{sum(1 for r in built if 'P' in r[3])} premium, "
          f"{sum(1 for r in built if r[3] == 'LP')} both)")
